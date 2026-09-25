"""Unit tests for Phase 2 Step 5: PDF + DOCX Regulatory Document Parsers.

Tests deterministic parsing, page tracking, heading extraction, table/list handling,
scanned PDF detection, factory dispatch, error handling, and RegulatoryChunker compatibility.
"""

import io
import pytest
import pypdf
import docx

from app.models.document import RegulatoryDocument, RegulatoryProvenance
from app.services.chunking.regulatory_chunker import RegulatoryChunker
from app.services.parsers import (
    DocxDocumentParser,
    PdfDocumentParser,
    get_parser,
    parse_regulatory_document,
)


# ============================================================================
# HELPER FIXTURE BUILDERS (Programmatic in-memory fixtures)
# ============================================================================

def make_test_pdf_bytes(pages_text: list) -> bytes:
    """Build a minimal valid multi-page PDF in-memory with text streams and exact xref."""
    font_obj_num = 3
    obj_num = 4
    page_indices = []
    content_objs = []

    for p_idx, lines in enumerate(pages_text):
        page_obj_num = obj_num
        content_obj_num = obj_num + 1
        page_indices.append(page_obj_num)

        stream_lines = ["BT", "/F1 12 Tf", "72 720 Td"]
        for line in lines:
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            stream_lines.append(f"({escaped}) Tj")
            stream_lines.append("0 -18 Td")
        stream_lines.append("ET")
        stream_bytes = "\n".join(stream_lines).encode("latin-1")

        content_objs.append((content_obj_num, stream_bytes))
        obj_num += 2

    catalog_bytes = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids_str = " ".join(f"{k} 0 R" for k in page_indices)
    pages_bytes = f"<< /Type /Pages /Kids [{kids_str}] /Count {len(page_indices)} >>".encode("latin-1")
    font_bytes = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"

    all_objs = [(1, catalog_bytes), (2, pages_bytes), (3, font_bytes)]

    for page_idx, (content_num, stream_bytes) in zip(page_indices, content_objs):
        page_dict = f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_num} 0 R /Resources << /Font << /F1 3 0 R >> >> >>".encode("latin-1")
        all_objs.append((page_idx, page_dict))

        content_dict = f"<< /Length {len(stream_bytes)} >>\nstream\n".encode("latin-1") + stream_bytes + b"\nendstream"
        all_objs.append((content_num, content_dict))

    all_objs.sort(key=lambda x: x[0])

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    xref_offsets = [0]

    for num, body in all_objs:
        offset = out.tell()
        xref_offsets.append(offset)
        out.write(f"{num} 0 obj\n".encode("latin-1"))
        out.write(body)
        out.write(b"\nendobj\n")

    startxref = out.tell()
    out.write(f"xref\n0 {len(all_objs) + 1}\n".encode("latin-1"))
    out.write(b"0000000000 65535 f \n")
    for off in xref_offsets[1:]:
        out.write(f"{off:010d} 00000 n \n".encode("latin-1"))

    out.write(f"trailer\n<< /Size {len(all_objs) + 1} /Root 1 0 R >>\nstartxref\n{startxref}\n%%EOF".encode("latin-1"))
    return out.getvalue()


def make_blank_pdf_bytes(num_pages: int = 1) -> bytes:
    """Build a multi-page PDF with no text layer (simulating scanned / image-only document)."""
    writer = pypdf.PdfWriter()
    for _ in range(num_pages):
        writer.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def make_test_docx_bytes(
    title: str = "Mock Monograph",
    sections_data: list = None,
) -> bytes:
    """Build a valid DOCX document in-memory containing headings, lists, and tables."""
    doc = docx.Document()

    if sections_data is None:
        sections_data = [
            {
                "heading": "1. INDICATIONS AND USAGE",
                "level": 1,
                "paragraphs": ["Drug V is indicated for severe clinical symptoms."],
                "subsections": [
                    {
                        "heading": "1.1 Pediatric Population",
                        "level": 2,
                        "paragraphs": ["Not approved for pediatric patients."],
                        "bullets": ["Under 12 years: contraindicated.", "12-18 years: safety unestablished."],
                    }
                ],
            },
            {
                "heading": "2. DOSAGE AND ADMINISTRATION",
                "level": 1,
                "paragraphs": ["Administer in accordance with the following table:"],
                "table": [
                    ["Weight Bracket", "Daily Dosage"],
                    ["< 50 kg", "10 mg"],
                    [">= 50 kg", "20 mg"],
                ],
            },
        ]

    for sec in sections_data:
        doc.add_heading(sec["heading"], level=sec.get("level", 1))
        for p in sec.get("paragraphs", []):
            doc.add_paragraph(p)
        for b in sec.get("bullets", []):
            doc.add_paragraph(b, style="List Bullet")
        if "table" in sec:
            table_data = sec["table"]
            tbl = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
            for r_idx, row in enumerate(table_data):
                for c_idx, val in enumerate(row):
                    tbl.cell(r_idx, c_idx).text = val

        for sub in sec.get("subsections", []):
            doc.add_heading(sub["heading"], level=sub.get("level", 2))
            for p in sub.get("paragraphs", []):
                doc.add_paragraph(p)
            for b in sub.get("bullets", []):
                doc.add_paragraph(b, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ============================================================================
# 1. PDF PARSER TESTS
# ============================================================================

def test_pdf_parser_can_parse():
    """Verify PDF parser detection via filename, MIME type, and binary signature."""
    parser = PdfDocumentParser()
    assert parser.can_parse(b"%PDF-1.4...", filename="label.pdf") is True
    assert parser.can_parse(b"", filename="document.PDF") is True
    assert parser.can_parse(b"", mime_type="application/pdf") is True
    assert parser.can_parse(b"%PDF-1.7\nsome data") is True
    assert parser.can_parse(b"NOT A PDF", filename="test.txt") is False


def test_pdf_parser_multipage_text_and_traceability():
    """Verify PDF parser extracts multi-page text and preserves page breadcrumbs and section outlines."""
    pdf_bytes = make_test_pdf_bytes([
        ["1. INDICATIONS AND USAGE", "Drug Alpha is indicated for hypertension in adult patients."],
        ["2. DOSAGE AND ADMINISTRATION", "The recommended dosage is 10 mg taken orally once daily."],
    ])

    parser = PdfDocumentParser()
    doc = parser.parse(
        pdf_bytes,
        document_id="doc_pdf_001",
        document_name="Drug Alpha Label",
        metadata={"product_name": "Drug Alpha", "jurisdiction": "US_FDA"},
    )

    assert isinstance(doc, RegulatoryDocument)
    assert doc.document_id == "doc_pdf_001"
    assert doc.title == "Drug Alpha Label"
    assert doc.product_name == "Drug Alpha"
    assert "Pages: 2" in (doc.provenance.exact_location or "")
    assert len(doc.sections) == 2

    # Verify Page 1 section
    sec1 = doc.sections[0]
    assert "INDICATIONS" in sec1.heading_raw.upper()
    assert sec1.section_number == "1"
    assert "[Page 1]" in (sec1.raw_text or "")
    assert "hypertension" in (sec1.raw_text or "")

    # Verify Page 2 section
    sec2 = doc.sections[1]
    assert "DOSAGE" in sec2.heading_raw.upper()
    assert sec2.section_number == "2"
    assert "[Page 2]" in (sec2.raw_text or "")
    assert "10 mg taken orally" in (sec2.raw_text or "")


def test_pdf_parser_empty_and_corrupt():
    """Verify PDF parser handles empty content and invalid/corrupt bytes gracefully."""
    parser = PdfDocumentParser()

    empty_doc = parser.parse(b"", document_id="empty_pdf")
    assert len(empty_doc.sections) == 0
    assert "Pages: 0" in (empty_doc.provenance.exact_location or "")

    with pytest.raises(ValueError, match="Invalid PDF content"):
        parser.parse(b"%PDF-1.4 CORRUPTED DATA WITHOUT XREF")


def test_pdf_parser_scanned_image_only_detection():
    """Verify scanned / image-only PDF detection preserves page metadata without fabricating text."""
    blank_pdf = make_blank_pdf_bytes(num_pages=3)
    parser = PdfDocumentParser()
    doc = parser.parse(blank_pdf, document_id="doc_scanned_001")

    assert "Scanned" in (doc.provenance.exact_location or "")
    assert "Pages: 3" in (doc.provenance.exact_location or "")
    assert len(doc.sections) == 3

    for idx, sec in enumerate(doc.sections):
        page_num = idx + 1
        assert f"Page {page_num} (Scanned)" in sec.heading_raw
        assert "No digital text layer detected" in (sec.raw_text or "")


# ============================================================================
# 2. DOCX PARSER TESTS
# ============================================================================

def test_docx_parser_can_parse():
    """Verify DOCX parser detection via filename, MIME type, and PK signature."""
    parser = DocxDocumentParser()
    assert parser.can_parse(b"", filename="dossier.docx") is True
    assert parser.can_parse(b"", mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document") is True
    assert parser.can_parse(b"PK\x03\x04...word/document.xml...") is True
    assert parser.can_parse(b"NOT A DOCX", filename="dossier.txt") is False


def test_docx_parser_hierarchy_lists_and_tables():
    """Verify DOCX parser preserves heading hierarchy, lists, and tables with cell traceability."""
    docx_bytes = make_test_docx_bytes()
    parser = DocxDocumentParser()
    doc = parser.parse(
        docx_bytes,
        document_id="doc_docx_001",
        document_name="Product Dossier",
        metadata={"jurisdiction": "US_FDA"},
    )

    assert isinstance(doc, RegulatoryDocument)
    assert doc.document_id == "doc_docx_001"
    assert doc.title == "Product Dossier"
    assert len(doc.sections) == 2

    # Check Section 1 with subsection
    sec1 = doc.sections[0]
    assert "1. INDICATIONS" in sec1.heading_raw
    assert "Drug V is indicated" in (sec1.raw_text or "")
    assert len(sec1.subsections) == 1

    sub11 = sec1.subsections[0]
    assert "1.1 Pediatric Population" in sub11.heading_raw
    assert sub11.parent_section_id == sec1.section_id
    assert "* Under 12 years" in (sub11.raw_text or "")

    # Check Section 2 with embedded Table
    sec2 = doc.sections[1]
    assert "2. DOSAGE" in sec2.heading_raw
    assert "[Table 1: 3 rows x 2 cols]" in (sec2.raw_text or "")
    assert "| Weight Bracket | Daily Dosage |" in (sec2.raw_text or "")
    assert "| < 50 kg | 10 mg |" in (sec2.raw_text or "")


def test_docx_parser_empty_and_corrupt():
    """Verify DOCX parser handles empty input and corrupted bytes."""
    parser = DocxDocumentParser()

    empty_doc = parser.parse(b"", document_id="empty_docx")
    assert len(empty_doc.sections) == 0

    with pytest.raises(ValueError, match="Invalid DOCX content"):
        parser.parse(b"PK\x03\x04 CORRUPTED NOT A REAL ZIP")


# ============================================================================
# 3. FACTORY DISPATCH TESTS (PDF & DOCX)
# ============================================================================

def test_parser_factory_pdf_and_docx():
    """Verify get_parser selects PdfDocumentParser and DocxDocumentParser appropriately."""
    assert isinstance(get_parser(filename="label.pdf"), PdfDocumentParser)
    assert isinstance(get_parser(mime_type="application/pdf"), PdfDocumentParser)
    assert isinstance(get_parser(content=b"%PDF-1.4..."), PdfDocumentParser)

    assert isinstance(get_parser(filename="protocol.docx"), DocxDocumentParser)
    assert isinstance(get_parser(mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"), DocxDocumentParser)


# ============================================================================
# 4. REGULATORY CHUNKER COMPATIBILITY (PDF & DOCX)
# ============================================================================

def test_pdf_chunker_compatibility():
    """Verify that documents parsed from PDF flow smoothly into RegulatoryChunker."""
    pdf_bytes = make_test_pdf_bytes([
        [
            "1. INDICATIONS AND USAGE",
            "Drug Beta is indicated for chronic heart conditions.",
            "It must be taken with water.",
        ],
        [
            "2. DOSAGE AND ADMINISTRATION",
            "Initial dose is 5 mg once daily.",
            "Titrate up to 10 mg after 4 weeks.",
        ],
    ])

    doc = parse_regulatory_document(
        content=pdf_bytes,
        filename="heart_label.pdf",
        document_id="doc_pdf_chunk_test",
    )

    chunker = RegulatoryChunker()
    chunks = chunker.chunk_document(doc)

    assert len(chunks) >= 2
    for chk in chunks:
        assert chk.document_id == "doc_pdf_chunk_test"
        assert chk.chunk_id.startswith("chk_")
        assert len(chk.content.strip()) > 0
        # Verify conversion to backward-compatible content item
        item = chk.to_content_item()
        assert item.document_id == "doc_pdf_chunk_test"


def test_docx_chunker_compatibility():
    """Verify that documents parsed from DOCX flow smoothly into RegulatoryChunker."""
    docx_bytes = make_test_docx_bytes()

    doc = parse_regulatory_document(
        content=docx_bytes,
        filename="dossier.docx",
        document_id="doc_docx_chunk_test",
    )

    chunker = RegulatoryChunker()
    chunks = chunker.chunk_document(doc)

    assert len(chunks) >= 3
    # Check that table row chunks and paragraph chunks are properly generated
    chunk_types = {c.chunk_type for c in chunks}
    assert "paragraph" in chunk_types or "table_row" in chunk_types or "bullet" in chunk_types

    for chk in chunks:
        assert chk.document_id == "doc_docx_chunk_test"
        item = chk.to_content_item()
        assert item.document_id == "doc_docx_chunk_test"

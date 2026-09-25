"""DOCX regulatory document parser.

Parses Microsoft Word (.docx) regulatory dossiers, preserving document reading order,
heading hierarchy (Heading 1-6), paragraphs, bullet/numbered lists, and tables with
cell and row traceability.
"""

import io
import re
from typing import Any, Dict, List, Optional, Tuple, Union
from uuid import uuid4

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.models.document import (
    CanonicalSectionConcept,
    RegulatoryDocument,
    RegulatoryProvenance,
    RegulatorySection,
)
from app.services.chunking.outline_engine import HeadingInfo, RegulatoryOutlineEngine
from app.services.parsers.base_parser import BaseDocumentParser


class DocxDocumentParser(BaseDocumentParser):
    """Deterministic parser for Microsoft Word (.docx) regulatory documents."""

    def __init__(self):
        self.outline_engine = RegulatoryOutlineEngine()

    def can_parse(
        self,
        content: Union[str, bytes],
        filename: Optional[str] = None,
        mime_type: Optional[str] = None,
    ) -> bool:
        """Evaluate if content or filename indicates DOCX format."""
        if filename and filename.lower().endswith(".docx"):
            return True
        if mime_type and mime_type.lower() in (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/docx",
            "application/msword",
        ):
            return True

        # DOCX is a zip file starting with PK header and containing word/document.xml
        if isinstance(content, (bytes, bytearray)):
            return content.startswith(b"PK\x03\x04") and b"word/document.xml" in content[:4096]
        elif isinstance(content, str):
            raw_b = content.encode("latin-1", errors="ignore")
            return raw_b.startswith(b"PK\x03\x04") and b"word/document.xml" in raw_b[:4096]

        return False

    def parse(
        self,
        content: Union[str, bytes],
        document_id: Optional[str] = None,
        document_name: Optional[str] = None,
        provenance: Optional[RegulatoryProvenance] = None,
        metadata: Optional[Dict[str, Any]] = None,
    ) -> RegulatoryDocument:
        """Parse DOCX content into a hierarchical RegulatoryDocument container."""
        docx_bytes = self._to_bytes(content)

        meta = metadata or {}
        doc_id = document_id or meta.get("document_id") or f"doc_{uuid4().hex[:10]}"
        doc_title = document_name or meta.get("document_name") or meta.get("title")
        doc_type = meta.get("document_type", "REGULATORY_LABEL")
        jurisdiction = meta.get("jurisdiction", "US_FDA")
        product = meta.get("product_name") or meta.get("product")
        drug = meta.get("active_ingredient") or meta.get("drug")
        version = meta.get("version", "1.0")

        prov = provenance or self.default_provenance(
            source_repo=meta.get("source", "InternalDraft"),
            identifier=meta.get("source_identifier", doc_id),
            url=meta.get("source_url"),
        )

        if not docx_bytes or len(docx_bytes.strip()) == 0:
            prov.exact_location = "Paragraphs: 0, Tables: 0"
            return RegulatoryDocument(
                document_id=doc_id,
                title=doc_title or "Empty DOCX Document",
                document_type=doc_type,
                jurisdiction=jurisdiction,
                product_name=product,
                active_ingredient=drug,
                version=version,
                provenance=prov,
                sections=[],
                raw_content="",
            )

        try:
            doc = docx.Document(io.BytesIO(docx_bytes))
        except Exception as exc:
            raise ValueError(f"Invalid DOCX content: {exc}") from exc

        # Extract core document properties if title not provided
        try:
            core_props = doc.core_properties
            if not doc_title and core_props.title and core_props.title.strip():
                doc_title = core_props.title.strip()
        except Exception:
            pass

        if not doc_title:
            doc_title = "DOCX Regulatory Document"

        # Traverse document elements in exact source order
        raw_sections_data: List[Dict[str, Any]] = []
        current_heading_info: Optional[HeadingInfo] = None
        current_heading_level: int = 1
        current_body_lines: List[str] = []
        order_idx = 0
        table_count = 0
        paragraph_count = 0

        for child in doc.element.body:
            tag = child.tag.split("}")[-1]

            if tag == "p":
                paragraph_count += 1
                p = Paragraph(child, doc)
                p_text = p.text.strip()
                if not p_text:
                    continue

                style_name = p.style.name if p.style else ""
                is_heading, level = self._detect_heading(p_text, style_name, order_idx)

                if is_heading:
                    # Flush previous section
                    if current_heading_info or current_body_lines:
                        sec_dict = self._package_raw_section(
                            heading_info=current_heading_info,
                            level=current_heading_level,
                            body_lines=current_body_lines,
                            order_index=order_idx,
                        )
                        raw_sections_data.append(sec_dict)
                        order_idx += 1
                        current_body_lines = []

                    heading_info = self.outline_engine.parse_heading(p_text, order_index=order_idx)
                    if not heading_info:
                        heading_info = HeadingInfo(
                            heading_raw=p_text,
                            heading_normalized=self.outline_engine.normalize_concept(p_text),
                            level=level,
                            order_index=order_idx,
                        )

                    current_heading_info = heading_info
                    current_heading_level = level
                else:
                    # Check if styled as a list item
                    formatted_line = self._format_paragraph_text(p_text, style_name)
                    current_body_lines.append(formatted_line)

            elif tag == "tbl":
                table_count += 1
                tbl = Table(child, doc)
                md_table = self._table_to_markdown(tbl, table_index=table_count)
                if md_table:
                    current_body_lines.append("\n" + md_table + "\n")

        # Flush final section
        if current_heading_info or current_body_lines:
            sec_dict = self._package_raw_section(
                heading_info=current_heading_info,
                level=current_heading_level,
                body_lines=current_body_lines,
                order_index=order_idx,
            )
            raw_sections_data.append(sec_dict)

        # Build hierarchical sections tree
        root_sections = self._build_hierarchical_sections(raw_sections_data, doc_id)

        # If title is still generic and first section is level 1, infer title
        if doc_title == "DOCX Regulatory Document" and raw_sections_data and raw_sections_data[0]["level"] == 1:
            doc_title = raw_sections_data[0]["heading_raw"]

        # Reconstruct raw content for traceability
        all_blocks = []
        for s in raw_sections_data:
            block = f"## {s['heading_raw']}\n{s['raw_text']}"
            all_blocks.append(block)
        combined_raw = "\n\n".join(all_blocks)

        prov.exact_location = f"Paragraphs: {paragraph_count}, Tables: {table_count}"

        return RegulatoryDocument(
            document_id=doc_id,
            title=doc_title,
            document_type=doc_type,
            jurisdiction=jurisdiction,
            product_name=product,
            active_ingredient=drug,
            version=version,
            provenance=prov,
            sections=root_sections,
            raw_content=combined_raw,
        )

    def _detect_heading(self, text: str, style_name: str, order_idx: int) -> Tuple[bool, int]:
        """Detect if paragraph is a heading by style name or numbered outline pattern."""
        # 1. Check style name
        if style_name:
            style_lower = style_name.lower()
            if "heading 1" in style_lower or style_lower == "title":
                return True, 1
            elif "heading 2" in style_lower:
                return True, 2
            elif "heading 3" in style_lower:
                return True, 3
            elif "heading 4" in style_lower:
                return True, 4
            elif "heading 5" in style_lower:
                return True, 5
            elif "heading 6" in style_lower:
                return True, 6

        # 2. Check outline numbering pattern
        heading_info = self.outline_engine.parse_heading(text, order_index=order_idx)
        if heading_info and heading_info.section_number:
            level = heading_info.level or 1
            return True, level

        return False, 1

    @staticmethod
    def _format_paragraph_text(text: str, style_name: str) -> str:
        """Format paragraph preserving list marker semantics for downstream extraction."""
        style_lower = style_name.lower()
        if "list bullet" in style_lower:
            if not text.startswith(("*", "-", "•")):
                return f"* {text}"
        elif "list number" in style_lower:
            if not re.match(r'^\d+[\.\)]\s+', text):
                return f"1. {text}"
        return text

    @staticmethod
    def _table_to_markdown(tbl: Table, table_index: int) -> str:
        """Convert a python-docx Table into a Markdown table with location metadata."""
        rows_data: List[List[str]] = []
        for row in tbl.rows:
            cell_texts = []
            for cell in row.cells:
                # Clean embedded newlines within table cells
                clean_cell = " ".join(cell.text.split())
                cell_texts.append(clean_cell)
            if any(cell_texts):
                rows_data.append(cell_texts)

        if not rows_data:
            return ""

        max_cols = max(len(r) for r in rows_data)
        if max_cols == 0:
            return ""

        for r in rows_data:
            while len(r) < max_cols:
                r.append("")

        num_rows = len(rows_data)
        meta_comment = f"[Table {table_index}: {num_rows} rows x {max_cols} cols]"

        header = rows_data[0]
        md_lines = [meta_comment, "| " + " | ".join(header) + " |"]
        md_lines.append("| " + " | ".join(["---"] * max_cols) + " |")
        for row in rows_data[1:]:
            md_lines.append("| " + " | ".join(row) + " |")

        return "\n".join(md_lines)

    @staticmethod
    def _package_raw_section(
        heading_info: Optional[HeadingInfo],
        level: int,
        body_lines: List[str],
        order_index: int,
    ) -> Dict[str, Any]:
        """Bundle section details into an intermediate dictionary."""
        title = heading_info.heading_raw if heading_info else "General Content"
        sec_num = heading_info.section_number if heading_info else None
        norm = heading_info.heading_normalized if heading_info else None
        body = "\n\n".join(b for b in body_lines if b).strip()

        return {
            "heading_raw": title,
            "section_number": sec_num,
            "heading_normalized": norm,
            "level": level,
            "raw_text": body,
            "order_index": order_index,
        }

    def _build_hierarchical_sections(
        self,
        raw_sections_data: List[Dict[str, Any]],
        doc_id: str,
    ) -> List[RegulatorySection]:
        """Convert list of section dictionaries into nested RegulatorySection hierarchy."""
        if not raw_sections_data:
            return []

        root_sections: List[RegulatorySection] = []
        stack: List[Tuple[int, RegulatorySection]] = []

        for data in raw_sections_data:
            level = data["level"]
            sec_num = data["section_number"]
            title = data["heading_raw"]
            sec_slug = re.sub(r'[^a-zA-Z0-9]', '_', (sec_num or title).lower())[:20]
            sec_id = f"sec_{doc_id}_{data['order_index']}_{sec_slug}"

            section = RegulatorySection(
                section_id=sec_id,
                section_number=sec_num,
                heading_raw=title,
                heading_normalized=data["heading_normalized"],
                order_index=data["order_index"],
                raw_text=data["raw_text"],
                subsections=[],
            )

            while stack and stack[-1][0] >= level:
                stack.pop()

            if stack:
                parent_sec = stack[-1][1]
                section.parent_section_id = parent_sec.section_id
                parent_sec.subsections.append(section)
            else:
                root_sections.append(section)

            stack.append((level, section))

        return root_sections

    @staticmethod
    def _to_bytes(content: Union[str, bytes]) -> bytes:
        """Safely convert string or bytes input to raw bytes."""
        if isinstance(content, (bytes, bytearray)):
            return bytes(content)
        if isinstance(content, str):
            try:
                return content.encode("latin-1")
            except Exception:
                return content.encode("utf-8", errors="replace")
        return bytes(content)
